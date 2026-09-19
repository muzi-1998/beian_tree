from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

import openpyxl


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"
INTERMEDIATE_DIR = OUTPUT_DIR / "intermediate"

EXPECTED = {
    "SUMO_process_inputs_1min.xlsx": {"Raw": (524161, 28), "Clean": (524161, 28), "Derived_m3d": (524161, 17), "QA": (28, 14), "DataDictionary": (31, 8)},
    "SUMO_process_inputs_5min.xlsx": {"Raw": (104833, 28), "Clean": (104833, 28), "Derived_m3d": (104833, 17), "QA": (28, 14), "DataDictionary": (31, 8)},
    "SUMO_process_inputs_1h.xlsx": {"Raw": (8737, 28), "Clean": (8737, 28), "Derived_m3d": (8737, 17), "QA": (28, 14), "DataDictionary": (31, 8)},
    "SUMO_DO_ORP_1min.xlsx": {"Raw": (524161, 15), "Clean": (524161, 15), "QA": (15, 14), "DataDictionary": (15, 8)},
    "SUMO_DO_ORP_5min.xlsx": {"Raw": (104833, 15), "Clean": (104833, 15), "QA": (15, 14), "DataDictionary": (15, 8)},
}


def structural_checks() -> dict[str, object]:
    report: dict[str, object] = {}
    for filename, sheets in EXPECTED.items():
        path = OUTPUT_DIR / filename
        with ZipFile(path) as archive:
            bad_member = archive.testzip()
            panes = []
            for index in range(1, len(sheets) + 1):
                xml = archive.read(f"xl/worksheets/sheet{index}.xml")[:12000].decode("utf-8", "ignore")
                pane = re.search(r"<pane[^>]*/>", xml)
                panes.append(pane.group(0) if pane else "")

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        if workbook.sheetnames != list(sheets):
            raise AssertionError((filename, workbook.sheetnames, list(sheets)))
        actual = {}
        for index, (name, shape) in enumerate(sheets.items()):
            worksheet = workbook[name]
            header = [
                cell.value
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))[: min(5, worksheet.max_column)]
            ]
            if (worksheet.max_row, worksheet.max_column) != shape:
                raise AssertionError((filename, name, (worksheet.max_row, worksheet.max_column), shape))
            if 'state="frozen"' not in panes[index]:
                raise AssertionError((filename, name, panes[index]))
            if worksheet["A1"].font.name != "Arial" or not worksheet["A1"].font.bold:
                raise AssertionError((filename, name, "header style"))
            actual[name] = {
                "rows": worksheet.max_row,
                "cols": worksheet.max_column,
                "pane": panes[index],
                "header": header,
                "font": worksheet["A1"].font.name,
                "bold": worksheet["A1"].font.bold,
                "fill": worksheet["A1"].fill.fgColor.rgb,
            }
        workbook.close()
        report[filename] = {"bytes": path.stat().st_size, "zip_test": bad_member, "sheets": actual}
    return report


def _sample_rows(path: Path, targets: set[int]) -> dict[int, dict[str, str]]:
    samples = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for index, row in enumerate(reader, start=1):
            if index in targets:
                samples[index] = row
    return samples


def formula_checks() -> dict[str, object]:
    checks = {}
    for scale, rows in (("1min", 524160), ("5min", 104832), ("1h", 8736)):
        targets = {1, rows // 2, rows}
        clean = _sample_rows(INTERMEDIATE_DIR / f"process_{scale}_clean.csv", targets)
        derived = _sample_rows(INTERMEDIATE_DIR / f"process_{scale}_derived.csv", targets)
        results = []
        for index in sorted(targets):
            source, result = clean[index], derived[index]

            def number(value: str) -> float:
                return float(value) if value else math.nan

            def close(left: float, right: float) -> bool:
                return (math.isnan(left) and math.isnan(right)) or math.isclose(left, right, rel_tol=1e-7, abs_tol=1e-6)

            expected_mbr = 24 * (
                number(source["剩余污泥流量（m³/h）"])
                + number(source["2#生物池外回流流量（m³/h）"])
                + number(source["1#生物池外回流流量（m³/h）"])
            )
            total_n = number(source["进水总氮（mg/L）"])
            expected_ratio = number(source["进水氨氮（mg/L）"]) / total_n if total_n > 0 else math.nan
            assertions = {
                "flow_x24": close(number(result["进水瞬时流量（m³/d）"]), 24 * number(source["进水瞬时流量（m³/h）"])),
                "mbr": close(number(result["MBR循环流（m³/d）"]), expected_mbr),
                "ratio": close(number(result["frSNHx_TKN（氨氮/总氮）"]), expected_ratio),
                "cod_proxy": close(number(result["frSU_SCCOD代理量（mg/L）"]), number(source["出水COD（mg/L）"]) - 2.5),
            }
            if not all(assertions.values()):
                raise AssertionError((scale, index, assertions))
            results.append({"row": index, "timestamp": source["日期"], **{key: "pass" for key in assertions}})
        checks[scale] = results
    return checks


def main() -> None:
    structural = structural_checks()
    formulas = formula_checks()
    (OUTPUT_DIR / "workbook_structural_QA.json").write_text(
        json.dumps(structural, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUTPUT_DIR / "derived_formula_QA.json").write_text(
        json.dumps(formulas, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("5 workbook structures and 12 derived-value checks passed")


if __name__ == "__main__":
    main()

