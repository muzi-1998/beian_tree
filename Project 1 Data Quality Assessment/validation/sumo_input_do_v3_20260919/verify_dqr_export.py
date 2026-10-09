"""Independently reconcile the delivered forcing files and audit workbook."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from export_dqr_forcing import DO, DO6, ORP, HERE, OUT, sha256


def csv(name):
    return pd.read_csv(OUT / name, index_col="datetime", parse_dates=["datetime"])


def verify(verify_hashes=False):
    manifest = json.loads((OUT / "DQR_forcing_manifest.json").read_text(encoding="utf-8"))
    contract = manifest["contract"]
    checks = {}

    def check(name, condition):
        checks[name] = bool(condition)
        if not condition:
            raise AssertionError(name)

    do = csv("DQR_DO8_5min_20260414_20260730.csv")
    orp = csv("DQR_ORP6_5min_20260414_20260730.csv")
    start, end = pd.Timestamp(contract["validation_start"]), pd.Timestamp(contract["end_exclusive"])
    expected = pd.date_range(start, end, freq="5min", inclusive="left", name="datetime")
    check("DO_108d_timebase", do.index.equals(expected) and len(do) == 31104)
    check("ORP_108d_timebase", orp.index.equals(expected))
    check("DO_channel_order", do.columns.tolist() == ["time_days"] + DO)
    check("ORP_channel_order", orp.columns.tolist() == ["time_days"] + ORP)
    check("elapsed_days", np.allclose(do.time_days, (expected - start).total_seconds() / 86400))
    check("nonnegative_DO", (do[DO].stack() >= 0).all())
    check("negative_ORP_retained", (orp[ORP].min() < 0).all())

    pilot = manifest["selected_pilot"]
    first, last = pd.Timestamp(pilot["start"]), pd.Timestamp(pilot["end_exclusive"])
    pilot_index = pd.date_range(first, last, freq="5min", inclusive="left", name="datetime")
    for name, channels in [("pilot28d_DO8_DQR.csv", DO), ("pilot28d_DO6_Raw.csv", DO6), ("pilot28d_ORP6.csv", ORP)]:
        frame = csv(name)
        check(name + "_timebase", frame.index.equals(pilot_index) and len(frame) == 8064)
        check(name + "_complete", np.isfinite(frame[channels]).all().all())
    process = csv("pilot28d_common_process_inputs.csv")
    check("pilot_process_timebase", process.index.equals(pilot_index))
    check("pilot_12_selected_process_inputs_finite", len(process.columns) == 13 and np.isfinite(process).all().all())
    check("pilot_DO6_Raw_DQR_identical", np.array_equal(csv("pilot28d_DO8_DQR.csv")[DO6], csv("pilot28d_DO6_Raw.csv")[DO6]))

    audit = pd.read_parquet(OUT / "do_orp_strategy_audit_1min.parquet")
    check("audit_unique_sensor_minute", not audit.duplicated(["sensor_id", "datetime"]).any())
    known = audit.decision_at.notna()
    age = (audit.loc[known, "datetime"] - audit.loc[known, "decision_at"]).dt.total_seconds()
    check("no_future_or_expired_scores", (age.ge(0) & age.lt(60 * contract["score_expiry_minutes"])).all())
    unrecoverable = audit.action.eq("UNRECOVERABLE")
    check("unrecoverable_not_fabricated", audit.loc[unrecoverable, "forcing"].isna().all())
    kept = audit.action.eq("KEEP") & ~audit.zero_equivalent
    check("kept_measurements_unchanged", np.array_equal(audit.loc[kept, "forcing"], audit.loc[kept, "raw"], equal_nan=True))
    zero = audit.zero_equivalent & audit.action.eq("KEEP")
    check("zero_equivalence_only_DO4", audit.loc[zero, "sensor_id"].isin(["DO_1_4", "DO_2_4"]).all())
    check("zero_equivalence_raw_retained", audit.loc[zero, "forcing"].eq(0).all() and np.array_equal(audit.loc[zero, "raw"], audit.loc[zero, "dqr_measurement"]))
    do4 = audit.sensor_id.isin(["DO_1_4", "DO_2_4"])
    check("DO4_common_across_strategies_this_export", np.array_equal(audit.loc[do4, "common"], audit.loc[do4, "forcing"], equal_nan=True))

    summary = pd.read_csv(OUT / "channel_summary.csv").set_index("sensor_id")
    holdout = audit[audit.datetime.ge(start)]
    for sensor, frame in holdout.groupby("sensor_id", observed=True):
        line = summary.loc[sensor]
        check(sensor + "_summary_reconciles", len(frame) == line.minutes
              and int(frame.low_trigger.sum()) == line.low_trigger_minutes
              and int(frame.action.eq("UNRECOVERABLE").sum()) == line.unrecoverable_minutes
              and int(frame.zero_equivalent.sum()) == line.zero_equivalent_minutes)
        series = frame.set_index("datetime").forcing.sort_index()
        grouped = series.resample("5min")
        derived = grouped.mean().where(grouped.count().eq(5))
        delivered = (do if sensor in DO else orp)[sensor]
        check(sensor + "_5min_reconciles", derived.index.equals(delivered.index)
              and np.allclose(derived, delivered, rtol=1e-8, atol=1e-9, equal_nan=True))

    workbook = load_workbook(OUT / "DQR_DO_ORP_export_audit.xlsx", read_only=True, data_only=True)
    check("workbook_sheets", workbook.sheetnames == ["Channels", "Months", "Windows", "ProcessCoverage", "Contract"])
    windows = pd.read_csv(OUT / "continuous_input_windows.csv", parse_dates=["start", "end_exclusive"])
    excel_rows = list(workbook["Windows"].iter_rows(min_row=2, values_only=True))
    check("workbook_window_count", len(excel_rows) == len(windows))
    check("workbook_local_dates_exact", all(pd.Timestamp(row[1]) == source.start and pd.Timestamp(row[2]) == source.end_exclusive
          for row, source in zip(excel_rows, windows.itertuples())))
    excel_channels = list(workbook["Channels"].iter_rows(min_row=2, values_only=True))
    check("workbook_channel_values_exact", all(row[0] == sensor and np.allclose(row[1:], source.values, rtol=1e-12)
          for row, (sensor, source) in zip(excel_channels, summary.iterrows())) and len(excel_channels) == len(summary))
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!"}
    check("workbook_no_error_values", all(not isinstance(cell, str) or cell not in errors
          for sh in workbook for row in sh.iter_rows(values_only=True) for cell in row))
    workbook.close()

    if verify_hashes:
        for path, digest in manifest["input_sha256"].items():
            check("source_hash:" + Path(path).name, sha256(Path(path)) == digest)
        for name, digest in manifest["code_sha256"].items():
            check("code_hash:" + name, sha256(HERE / name) == digest)
        for name, digest in manifest["output_sha256"].items():
            check("output_hash:" + name, sha256(OUT / name) == digest)
    else:
        (OUT / "independent_export_QA.json").write_text(json.dumps(checks, indent=2), encoding="utf-8", newline="\n")
    print(f"{len(checks)} independent checks passed; frozen hashes checked: {verify_hashes}")
    return checks


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify-hashes", action="store_true")
    verify(parser.parse_args().verify_hashes)
